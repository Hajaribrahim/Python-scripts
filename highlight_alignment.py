import argparse
import re
import pandas as pd
import os
from openpyxl import load_workbook, Workbook
from openpyxl.styles import PatternFill

#help message
parser = argparse.ArgumentParser(
    description="""Used to highlight residues of interest in an aligned MSA file """,
    formatter_class=argparse.RawDescriptionHelpFormatter)
parser.add_argument("input_file", help="Input alignment file from Hagar's code (.csv or .xlsx) [REQUIRED]")
parser.add_argument("resi_list", help="Listed residues of interest to be highlighted (Format: column header: residues, Ex: 4HEA:M1) [REQUIRED]")
parser.add_argument("-o", "--output", help="Optional output XLSX filename")

args = parser.parse_args()

# Input and output 
input_file = args.input_file
resi_list = args.resi_list
output_file = args.output or (os.path.splitext(input_file)[0] + "_highlighted.xlsx")

print("Input file:", input_file)
print("Residue list:", resi_list)
print("Output file:", output_file)

#Read input resi list
highlight_resi = {}

with open(resi_list) as f:
    for line in f:
        line = line.strip()
        if not line or ":" not in line:
            continue
        key, values = line.split(":", 1)
        residues = [r.strip() for r in values.split(",")]
        highlight_resi[key.strip()] = residues

print("Parsed residue dictionary:")
for k, v in highlight_resi.items():
    print(f"{k}: {v}")

# Load input aligned file
ext = os.path.splitext(input_file)[1].lower()
if ext == ".xlsx":
    wb = load_workbook(input_file)
    ws = wb.active

elif ext == ".csv":
    df = pd.read_csv(input_file)

# Convert openpyxl worksheet
wb = Workbook()
ws = wb.active

for c_idx, header in enumerate(df.columns, start=1):
        ws.cell(row=1, column=c_idx, value=header)

for r_idx, row in enumerate(df.itertuples(index=False), start=2):
        for c_idx, value in enumerate(row, start=1):
            ws.cell(row=r_idx, column=c_idx, value=value)

# Converts "M1" = "1_M"
highlight_dict = {}
for org, res_list in highlight_resi.items():
    parsed = []
    for item in res_list:
        match = re.search(r"([A-Z])(\d+)", item)
        if match:
            aa = match.group(1)
            num = match.group(2)
            parsed.append(f"{num}_{aa}")
    highlight_dict[org] = parsed

start_color = {
    "yellow": "FFFFFF00",
    "red": "FFFF0000",
    "green": "FF00FF00",
    "blue": "FF0000FF",
    "orange": "FFFFA500",
    "purple": "FF800080",
    "pink": "FFFFC0CB",
    "cyan": "FF00FFFF",
    "magenta": "FFFF00FF",
    "black": "FF000000",
    "white": "FFFFFFFF",
    "gray": "FF808080"
}
highlight_color_name = "yellow" #Can be changed to any color above
fill_style = PatternFill(
    start_color=start_color[highlight_color_name],
    end_color=start_color[highlight_color_name],
    fill_type="solid")

#highlights cells
ws_headers = {}
for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=1), start=1):
    header = col_cells[0].value
    if header:
        ws_headers[col_idx] = str(header).strip().lower()

for col_idx, column_cells in enumerate(ws.iter_cols(min_row=2), start=1):
    header_clean = ws_headers.get(col_idx, "")
    for key_lower, org_key in {k.lower(): k for k in highlight_dict.keys()}.items():
        if key_lower in header_clean:  
            residues_to_highlight = highlight_dict[org_key]
            for cell in column_cells:
                if cell.value in residues_to_highlight:
                    cell.fill = fill_style
            break  

wb.save(output_file)
print(f"✅ Highlighted Excel file saved as {output_file}")
